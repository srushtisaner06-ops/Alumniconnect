import asyncio
import csv
import io
import json
import os
import re
from contextlib import asynccontextmanager
from datetime import timedelta
from typing import Optional

import openpyxl
from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pydantic import BaseModel, Field

from . import config, engine, importer, notifier, services
from .db import db, init
from .security import current_user, hash_password, make_jwt, verify_password

DEFAULT_EMAIL_SUBJECT = "Please update your career details – {{batch}} batch"
DEFAULT_EMAIL_BODY = ("Dear {{name}},\n\nThe department is collecting career information from the {{batch}} batch "
                      "for accreditation reports. It takes about two minutes.\n\nYour personal link:\n{{unique_link}}\n\n"
                      "Want reminders on Telegram instead? Open this link and press Start:\n{{telegram_link}}\n\n"
                      "Thank you,\nDepartment of Computer Science (AI)")
DEFAULT_TELEGRAM = "Hi {{name}}, please update your career details ({{batch}} batch): {{unique_link}}"


async def _worker():
    while True:
        await asyncio.sleep(config.WORKER_TICK_SECONDS)
        try:
            await asyncio.to_thread(services.run_all_running)
        except Exception as e:  # never let the worker die
            print("worker error:", e)


async def _telegram_poller():
    while True:
        await asyncio.sleep(config.TELEGRAM_POLL_SECONDS)
        try:
            await asyncio.to_thread(services.poll_telegram)
        except Exception as e:  # e.g. no network; keep trying
            print("telegram poll error:", e)


@asynccontextmanager
async def lifespan(app):
    init()
    with db() as c:
        if not c.execute("SELECT 1 FROM users").fetchone():
            c.execute("INSERT INTO users(email,password_hash) VALUES(?,?)", (config.ADMIN_EMAIL, hash_password(config.ADMIN_PASSWORD)))
        # First run only: load the real placement sheet through the same validated importer.
        if os.getenv("AUTOLOAD_REFERENCE", "1") == "1" and not c.execute("SELECT 1 FROM alumni").fetchone() \
                and config.REFERENCE_XLSX.exists():
            rep = importer.parse(config.REFERENCE_XLSX.name, config.REFERENCE_XLSX.read_bytes())
            if rep["ok"]:
                services.import_students(c, rep["students"], rep["graduation_year"], final_year_students=rep["final_year_students"])
    tasks = [asyncio.create_task(_worker()), asyncio.create_task(_telegram_poller())] if config.WORKER_ENABLED else []
    yield
    for t in tasks:
        t.cancel()


app = FastAPI(title="AlumniConnect", lifespan=lifespan)
app.add_middleware(CORSMiddleware, allow_origins=config.CORS_ORIGINS, allow_methods=["*"], allow_headers=["*"])


# ================================================================ auth
class Login(BaseModel):
    email: str
    password: str


@app.post("/api/auth/login")
def login(body: Login):
    with db() as c:
        u = c.execute("SELECT * FROM users WHERE email=?", (body.email.strip().lower(),)).fetchone()
    if not u or not verify_password(body.password, u["password_hash"]):
        raise HTTPException(401, "Incorrect email or password")
    return {"token": make_jwt(u["email"]), "email": u["email"]}


# ================================================================ dashboard / analytics
@app.get("/api/dashboard")
def dashboard(_=Depends(current_user)):
    with db() as c:
        data = services.analytics(c)
        recs = [dict(r) for r in c.execute("SELECT id,prn,status FROM alumni")]
        ids = engine.pending_list(recs)[:8]
        nxt = []
        for i in ids:
            r = c.execute("SELECT id,name,prn,email FROM alumni WHERE id=?", (i,)).fetchone()
            nxt.append(dict(r))
        data["next_pending"] = nxt
        data["campaigns"] = [dict(r) for r in c.execute("SELECT id,name,status FROM campaigns ORDER BY id DESC LIMIT 3")]
    return data


@app.get("/api/reports/analytics")
def reports(_=Depends(current_user)):
    with db() as c:
        return services.analytics(c)


@app.get("/api/reports/export")
def export(scope: str = "submissions", format: str = "xlsx", _=Depends(current_user)):
    if scope not in ("submissions", "placement") or format not in ("csv", "xlsx"):
        raise HTTPException(422, "scope must be submissions|placement, format csv|xlsx")
    with db() as c:
        meta = [dict(r) for r in c.execute("SELECT id,prn,graduation_year,status FROM alumni")]
        order = engine.sort_ids(meta, "graduation_year")
        rows = []
        for aid in order:
            a = services.alumnus_dict(c, c.execute("SELECT * FROM alumni WHERE id=?", (aid,)).fetchone())
            if scope == "placement":
                for o in a["offers"]:
                    rows.append([a["sr_no"], a["name"], a["prn"], a["class_div"], a["phone"], a["email"], o["employer"],
                                 o["package_lpa"], o["position"], o["letter_ref"], o["letter_link"]])
            else:
                d = a["details"] or {}
                rows.append([a["name"], a["prn"], a["graduation_year"], a["email"], a["phone"], a["status"],
                             a["career_status"] or "", a["submitted_at"] or "", a["reminders_sent"],
                             d.get("company") or d.get("university") or d.get("startup_name") or "",
                             d.get("designation") or d.get("degree") or d.get("business_domain") or "",
                             d.get("industry") or d.get("country") or d.get("stage") or "",
                             d.get("specialization") or d.get("notes") or ""])
    if scope == "placement":
        head = [h for _, h, _ in importer.COLUMNS]
    else:
        head = ["Name", "PRN", "Graduation Year", "Email", "Phone", "Status", "Career Status", "Submitted At", "Reminders Sent",
                "Company / University / Startup", "Designation / Degree / Domain", "Industry / Country / Stage", "Specialization / Notes"]
    fname = f"alumni_{scope}"
    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(head)
        w.writerows(rows)
        return Response(buf.getvalue().encode("utf-8-sig"), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="{fname}.csv"'})
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(head)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}.xlsx"'})


# ================================================================ alumni
@app.get("/api/alumni")
def list_alumni(search: str = "", status: str = "all", page: int = 1, page_size: int = Query(25, le=100), _=Depends(current_user)):
    q = search.strip()
    where, args = [], []
    with db() as c:
        if re.fullmatch(r"\d{8}", q):          # PRN lookup
            recs = [dict(r) for r in c.execute("SELECT id,prn FROM alumni")]
            hit = engine.search_prn(recs, q)
            where.append("a.id=?")
            args.append(hit if hit is not None else -1)
        elif q:
            like = f"%{q}%"
            where.append("(a.name LIKE ? OR a.prn LIKE ? OR a.email LIKE ? OR a.id IN (SELECT alumni_id FROM offers WHERE employer LIKE ?))")
            args += [like] * 4
        cutoff = services.overdue_cutoff(c)
        today = services.now().strftime("%Y-%m-%d")
        is_overdue = bool(cutoff and cutoff < today)
        if status == "COMPLETED":
            where.append("a.status='COMPLETED'")
        elif status == "OVERDUE":
            where.append("a.status='PENDING'" if is_overdue else "0")
        elif status == "PENDING":
            where.append("a.status='PENDING'" + (" AND 0" if is_overdue else ""))
        w = ("WHERE " + " AND ".join(where)) if where else ""
        total = c.execute(f"SELECT COUNT(*) n FROM alumni a {w}", args).fetchone()["n"]
        rows = c.execute(f"SELECT a.* FROM alumni a {w} ORDER BY a.sr_no, a.id LIMIT ? OFFSET ?",
                         args + [page_size, (max(page, 1) - 1) * page_size]).fetchall()
        items = []
        for r in rows:
            d = services.alumnus_dict(c, r)
            d.pop("token")
            items.append(d)
    return {"items": items, "total": total, "page": page, "page_size": page_size}


@app.get("/api/alumni/{aid}")
def get_alumnus(aid: int, _=Depends(current_user)):
    with db() as c:
        r = c.execute("SELECT * FROM alumni WHERE id=?", (aid,)).fetchone()
        if not r:
            raise HTTPException(404, "Not found")
        d = services.alumnus_dict(c, r)
        d["link"] = f"{config.PUBLIC_URL}/update?token={d.pop('token')}"
        d["history"] = [dict(x) for x in c.execute(
            "SELECT l.*, cm.name campaign FROM reminder_log l LEFT JOIN campaigns cm ON cm.id=l.campaign_id WHERE alumni_id=? ORDER BY l.id DESC LIMIT 20", (aid,))]
    return d


class AlumniPatch(BaseModel):
    telegram_chat_id: Optional[str] = Field(None, max_length=40)
    unsubscribed: Optional[bool] = None


@app.patch("/api/alumni/{aid}")
def patch_alumnus(aid: int, body: AlumniPatch, _=Depends(current_user)):
    with db() as c:
        r = c.execute("SELECT * FROM alumni WHERE id=?", (aid,)).fetchone()
        if not r:
            raise HTTPException(404, "Not found")
        if body.telegram_chat_id is not None:
            c.execute("UPDATE alumni SET telegram_chat_id=? WHERE id=?", (body.telegram_chat_id.strip() or None, aid))
        if body.unsubscribed is not None and int(body.unsubscribed) != r["unsubscribed"]:
            services.push_undo(c, "set_unsubscribed", {"id": aid, "previous": r["unsubscribed"]},
                               f"{'Unsubscribed' if body.unsubscribed else 'Re-subscribed'} {r['name']}")
            c.execute("UPDATE alumni SET unsubscribed=? WHERE id=?", (int(body.unsubscribed), aid))
    return {"ok": True}


@app.delete("/api/alumni/{aid}")
def delete_alumnus(aid: int, _=Depends(current_user)):
    with db() as c:
        r = c.execute("SELECT name FROM alumni WHERE id=?", (aid,)).fetchone()
        if not r:
            raise HTTPException(404, "Not found")
        services.push_undo(c, "delete_alumni", services.snapshot(c, aid), f"Deleted {r['name']}")
        c.execute("DELETE FROM alumni WHERE id=?", (aid,))
    return {"ok": True}


@app.get("/api/undo")
def undo_peek(_=Depends(current_user)):
    with db() as c:
        rows = c.execute("SELECT id,label FROM undo_actions ORDER BY id").fetchall()
        return {"count": len(rows), "next": rows[-1]["label"] if rows else None}


@app.post("/api/undo")
def undo(_=Depends(current_user)):
    with db() as c:
        label = services.undo_last(c)
    if label is None:
        raise HTTPException(404, "Nothing to undo")
    return {"undone": label}


# ================================================================ import
@app.get("/api/import/template")
def import_template(_=Depends(current_user)):
    return Response(importer.template_workbook(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="placement_template.xlsx"'})


MAX_UPLOAD = 5 * 1024 * 1024


async def _read_upload(file: UploadFile):
    data = await file.read()
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, "File too large (max 5 MB)")
    return data


def _public_report(rep):
    return {k: v for k, v in rep.items() if k != "students"}


@app.post("/api/import/validate")
async def import_validate(file: UploadFile = File(...), graduation_year: Optional[int] = Form(None), _=Depends(current_user)):
    rep = importer.parse(file.filename, await _read_upload(file), graduation_year)
    with db() as c:
        existing = {r["prn"] for r in c.execute("SELECT prn FROM alumni")}
    rep["already_in_database"] = sum(1 for s in rep["students"] if s["prn"] in existing)
    return _public_report(rep)


@app.post("/api/import")
async def import_commit(file: UploadFile = File(...), graduation_year: Optional[int] = Form(None),
                        replace: bool = Form(False), skip_invalid: bool = Form(False), _=Depends(current_user)):
    rep = importer.parse(file.filename, await _read_upload(file), graduation_year)
    fatal = rep["missing_columns"] or not rep["graduation_year"] or not rep["students"] or rep["header_row"] is None
    if fatal or (rep["errors"] and not skip_invalid):
        raise HTTPException(422, {"message": "File does not match the reference format", "report": _public_report(rep)})
    with db() as c:
        if replace:
            c.execute("DELETE FROM import_batches")   # their alumni are wiped below, so the file history goes too
        imp = c.execute("INSERT INTO import_batches(filename,graduation_year) VALUES(?,?)",
                        ((file.filename or "upload")[:200], rep["graduation_year"])).lastrowid
        res = services.import_students(c, rep["students"], rep["graduation_year"], replace=replace,
                                       final_year_students=rep["final_year_students"], import_id=imp)
        c.execute("UPDATE import_batches SET added=?, updated=?, unique_students=? WHERE id=?",
                  (res["added"], res["updated"], len(rep["students"]), imp))
    res["import_id"] = imp
    res["skipped_rows"] = len({e["row"] for e in rep["errors"] if e["row"]})
    return {"ok": True, **res, "unique_students": len(rep["students"]), "replaced": replace}


# ================================================================ campaigns
class CampaignIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)
    channel: str = "email"
    interval_hours: float = Field(72, ge=0.01, le=24 * 60)
    max_attempts: int = Field(3, ge=1, le=20)
    batch_size: int = Field(20, ge=1, le=500)
    deadline: Optional[str] = None
    email_subject: str = DEFAULT_EMAIL_SUBJECT
    email_body: str = DEFAULT_EMAIL_BODY
    telegram_body: str = DEFAULT_TELEGRAM
    # {"all": bool, "batches": [years], "import_ids": [ids], "alumni_ids": [ids]}.
    # Omitted (null) = leave unchanged on update / everyone on create (the behaviour before this feature).
    recipients: Optional[dict] = None


def _validate_campaign(b: CampaignIn):
    if b.channel not in ("email", "telegram", "both"):
        raise HTTPException(422, "channel must be email, telegram or both")
    bad = notifier.unknown_placeholders(b.email_subject, b.email_body, b.telegram_body)
    if bad:
        raise HTTPException(422, f"Unknown placeholder(s): {', '.join('{{' + x + '}}' for x in bad)}. "
                                 f"Allowed: {', '.join('{{' + x + '}}' for x in sorted(notifier.PLACEHOLDERS))}")
    used = ([b.email_body] if b.channel in ("email", "both") else []) + ([b.telegram_body] if b.channel in ("telegram", "both") else [])
    if any("{{unique_link}}" not in re.sub(r"\s", "", t) for t in used):
        raise HTTPException(422, "Every message template in use must include {{unique_link}} so alumni can open their form")
    if b.deadline and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", b.deadline):
        raise HTTPException(422, "deadline must be YYYY-MM-DD")


def _campaign_stats(c, row):
    d = dict(row)
    d["sent"] = c.execute("SELECT COUNT(*) n FROM reminder_log WHERE campaign_id=? AND status IN ('SENT','SIMULATED')", (row["id"],)).fetchone()["n"]
    d["failed"] = c.execute("SELECT COUNT(*) n FROM reminder_log WHERE campaign_id=? AND status='FAILED'", (row["id"],)).fetchone()["n"]
    d["eligible_now"] = len(services.eligible(c, row))
    d["reachable"] = services.eligible_total(c, row)
    sel = services.get_selection(c, row["id"])
    if sel is None:   # created before recipient selection existed: contacts everyone
        d["recipients"] = {"all": True, "batches": [], "import_ids": [], "alumni_ids": [], "people": [], "legacy": True}
        d["recipient_total"] = c.execute("SELECT COUNT(*) n FROM alumni").fetchone()["n"]
    else:
        d["recipients"] = {**sel, "people": services.people_for(c, sel["alumni_ids"])}
        d["recipient_total"] = len(services.resolve_selection(c, sel))
    return d


@app.get("/api/campaigns")
def list_campaigns(_=Depends(current_user)):
    with db() as c:
        rows = c.execute("SELECT * FROM campaigns ORDER BY id DESC").fetchall()
        return {"items": [_campaign_stats(c, r) for r in rows],
                "mode": {"email": "live" if notifier.email_configured() else "simulated",
                         "telegram": "live" if notifier.telegram_configured() else "simulated"},
                "placeholders": sorted(notifier.PLACEHOLDERS)}


@app.post("/api/campaigns")
def create_campaign(b: CampaignIn, _=Depends(current_user)):
    _validate_campaign(b)
    deadline = b.deadline or (services.now() + timedelta(days=14)).strftime("%Y-%m-%d")
    with db() as c:
        cid = c.execute("""INSERT INTO campaigns(name,channel,interval_hours,max_attempts,batch_size,deadline,email_subject,email_body,telegram_body)
                           VALUES(?,?,?,?,?,?,?,?,?)""", (b.name, b.channel, b.interval_hours, b.max_attempts, b.batch_size, deadline,
                                                           b.email_subject, b.email_body, b.telegram_body)).lastrowid
        if b.recipients is not None:
            services.save_selection(c, cid, services.normalize_selection(b.recipients))
    return {"id": cid}


@app.put("/api/campaigns/{cid}")
def update_campaign(cid: int, b: CampaignIn, _=Depends(current_user)):
    _validate_campaign(b)
    with db() as c:
        if not c.execute("SELECT 1 FROM campaigns WHERE id=?", (cid,)).fetchone():
            raise HTTPException(404, "Not found")
        c.execute("""UPDATE campaigns SET name=?,channel=?,interval_hours=?,max_attempts=?,batch_size=?,deadline=?,
                     email_subject=?,email_body=?,telegram_body=? WHERE id=?""",
                  (b.name, b.channel, b.interval_hours, b.max_attempts, b.batch_size, b.deadline, b.email_subject,
                   b.email_body, b.telegram_body, cid))
        if b.recipients is not None:
            services.save_selection(c, cid, services.normalize_selection(b.recipients))
    return {"ok": True}


# ---- recipient selection helpers (used by the campaign editor)
@app.get("/api/recipients/options")
def recipient_options(_=Depends(current_user)):
    with db() as c:
        return services.recipient_options(c)


@app.get("/api/recipients/search")
def recipient_search(q: str = "", _=Depends(current_user)):
    q = q.strip()
    if len(q) < 2:
        return {"items": []}
    like = f"%{q}%"
    with db() as c:
        rows = c.execute("SELECT id,name,prn,email,graduation_year,status FROM alumni "
                         "WHERE name LIKE ? OR prn LIKE ? OR email LIKE ? ORDER BY name LIMIT 15", (like, like, like)).fetchall()
    return {"items": [dict(r) for r in rows]}


class PreviewIn(BaseModel):
    recipients: dict
    channel: str = "email"
    max_attempts: int = Field(3, ge=1, le=20)


@app.post("/api/recipients/preview")
def recipient_preview(b: PreviewIn, _=Depends(current_user)):
    if b.channel not in ("email", "telegram", "both"):
        raise HTTPException(422, "channel must be email, telegram or both")
    with db() as c:
        return services.preview_recipients(c, services.normalize_selection(b.recipients), b.channel, b.max_attempts)


@app.post("/api/campaigns/{cid}/{action}")
def campaign_action(cid: int, action: str, _=Depends(current_user)):
    with db() as c:
        row = c.execute("SELECT * FROM campaigns WHERE id=?", (cid,)).fetchone()
        if not row:
            raise HTTPException(404, "Not found")
        st = row["status"]
        if action in ("start", "resume"):
            if st not in ("DRAFT", "PAUSED", "COMPLETED"):
                raise HTTPException(409, f"This round is already {st}")
            c.execute("UPDATE campaigns SET status='RUNNING', started_at=COALESCE(started_at, datetime('now')) WHERE id=?", (cid,))
        elif action == "pause":
            if st != "RUNNING":
                raise HTTPException(409, "Only a running round can be paused")
            c.execute("UPDATE campaigns SET status='PAUSED' WHERE id=?", (cid,))
        elif action == "run-now":
            if st != "RUNNING":
                raise HTTPException(409, "Start the round first")
        else:
            raise HTTPException(404, "Unknown action")
    return services.run_campaign(cid) if action == "run-now" else {"ok": True}


@app.get("/api/reminders")
def reminders(campaign_id: Optional[int] = None, limit: int = Query(100, le=500), _=Depends(current_user)):
    with db() as c:
        w, a = ("WHERE l.campaign_id=?", [campaign_id]) if campaign_id else ("", [])
        rows = c.execute(f"""SELECT l.*, al.name, al.prn, cm.name campaign FROM reminder_log l
                             LEFT JOIN alumni al ON al.id=l.alumni_id LEFT JOIN campaigns cm ON cm.id=l.campaign_id
                             {w} ORDER BY l.id DESC LIMIT ?""", a + [limit]).fetchall()
        return {"items": [dict(r) for r in rows]}


# ================================================================ public alumni portal (token only)
def _by_token(c, token):
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,64}", token or ""):
        raise HTTPException(404, "This link is not valid")
    r = c.execute("SELECT * FROM alumni WHERE token=?", (token,)).fetchone()
    if not r:
        raise HTTPException(404, "This link is not valid")
    return r


@app.get("/api/public/form")
def public_form(token: str):
    with db() as c:
        r = _by_token(c, token)
        best = c.execute("SELECT employer,position,package_lpa FROM offers WHERE alumni_id=? ORDER BY package_lpa DESC LIMIT 1", (r["id"],)).fetchone()
        return {"name": r["name"], "graduation_year": r["graduation_year"], "prn": r["prn"], "status": r["status"],
                "career_status": r["career_status"], "details": json.loads(r["details"]) if r["details"] else None,
                "prefill": {"company": best["employer"], "designation": best["position"]} if best else None}


class Submission(BaseModel):
    token: str
    career_status: str
    details: dict


REQUIRED = {"EMPLOYMENT": ["company", "designation", "industry"],
            "HIGHER_STUDIES": ["university", "country", "degree", "specialization"],
            "ENTREPRENEURSHIP": ["startup_name", "business_domain", "stage"]}
OPTIONAL = {"EMPLOYMENT": ["notes"], "HIGHER_STUDIES": [], "ENTREPRENEURSHIP": []}
STAGES = ["Idea", "Prototype / MVP", "Early revenue", "Growth"]


@app.post("/api/public/submit")
def public_submit(b: Submission):
    if b.career_status not in REQUIRED:
        raise HTTPException(422, "Choose Employment, Higher Studies or Entrepreneurship")
    clean = {}
    for k in REQUIRED[b.career_status]:
        v = str(b.details.get(k, "")).strip()
        if not v:
            raise HTTPException(422, f"'{k.replace('_', ' ')}' is required")
        if len(v) > 200:
            raise HTTPException(422, f"'{k.replace('_', ' ')}' is too long")
        clean[k] = v
    for k in OPTIONAL[b.career_status]:
        v = str(b.details.get(k, "")).strip()[:1000]
        if v:
            clean[k] = v
    if b.career_status == "ENTREPRENEURSHIP" and clean["stage"] not in STAGES:
        raise HTTPException(422, "Invalid stage")
    with db() as c:
        r = _by_token(c, b.token)
        again = r["status"] == "COMPLETED"
        c.execute("UPDATE alumni SET status='COMPLETED', career_status=?, details=?, submitted_at=datetime('now') WHERE id=?",
                  (b.career_status, json.dumps(clean), r["id"]))
    return {"ok": True, "already_submitted": again}
